"""组织架构（需求部门）管理：读取层级树 / Excel 批量导入 / 恢复默认 / 删除节点。

数据源为数据库 departments 表（见 services/departments.py）。
"""
from __future__ import annotations

import logging
from io import BytesIO
from pathlib import Path
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException, Response, UploadFile
from sqlalchemy.orm import Session

from ..database import get_db
from ..schemas import DepartmentProvinceUpdate
from ..services import departments as dept

logger = logging.getLogger(__name__)
router = APIRouter(prefix="/api/departments", tags=["departments"])

XLSX_SUFFIX = {".xlsx", ".xlsm"}


# ---------- Excel 模板 / 解析 ----------

def _build_template_bytes() -> bytes:
    """标准部门导入模板：每行填一个部门节点的完整路径 + 所在省份。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    headers = [
        "部门级联路径（从根到本级，用 / 分隔）*",
        "省份（如：广东；留空则继承上级）",
        "备注 / 说明（可选）",
    ]
    widths = [56, 30, 40]
    fill = PatternFill("solid", fgColor="DDEBF7")
    font = Font(bold=True)

    def style_sheet(ws) -> None:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for cell in ws[1]:
            cell.fill = fill
            cell.font = font
            cell.alignment = Alignment(vertical="center")

    wb = Workbook()
    ws = wb.active
    ws.title = "导入数据"
    ws.append(headers)
    style_sheet(ws)

    ws2 = wb.create_sheet("填写示例")
    ws2.append(headers)
    style_sheet(ws2)
    for p in [
        ["集团总部", "北京", "顶层节点：父级留空即根，省份填在本级"],
        ["集团总部/发展策划部", "", "第二级留空 = 继承上级（北京）"],
        ["集团总部/发展策划部/规划处", "", "第三级处室，同样继承北京"],
        ["省级电力公司", "广东", "另一个顶层：省份可与总部不同"],
        ["省级电力公司/调度中心", "", "继承广东"],
    ]:
        ws2.append(p)

    ws3 = wb.create_sheet("填写说明")
    note_lines = [
        "填写说明：",
        "1. 每行代表一个部门节点，A 列填写该节点从根到本级的完整路径，层级之间用 / 分隔。",
        "2. B 列填该节点所在省份；留空表示跟随上级，无需逐级重复填写。",
        "3. 岗位录入时省份直接由所选需求部门推导，因此组织架构里维护一次即可。",
        "4. 只需把希望存在的节点各写一行即可，父级顺序无要求——父级路径尚不存在会自动创建。",
        "5. 同一文件重复导入不会产生重复节点（按「同级名称唯一」去重）；",
        "   已存在节点若已设省份，不会被 Excel 中的空值覆盖。",
        "6. C 列为备注，导入时忽略，仅作人工说明。",
    ]
    for line in note_lines:
        ws3.append([line])
    ws3.column_dimensions["A"].width = 100

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _read_import_rows(file) -> list[dict]:
    """读取第一个工作表：第 1 列为「部门级联路径」，「省份」列可选，跳过表头与空行。"""
    from openpyxl import load_workbook

    try:
        wb = load_workbook(file, data_only=True, read_only=True)
    except Exception as e:
        raise HTTPException(400, f"无法解析 Excel 文件：{e}") from e

    ws = wb.worksheets[0]
    rows: list[dict] = []
    idx_prov: int | None = None
    for n, row in enumerate(ws.iter_rows(values_only=True), start=1):
        cells = list(row or ())
        text = str(cells[0]).strip() if cells and cells[0] is not None else ""
        if not text:
            continue
        if n == 1 and ("路径" in text or "部门" in text):  # 跳过表头行，并记下省份列位置
            idx_prov = next(
                (i for i, h in enumerate(cells)
                 if h is not None and ("省份" in str(h) or "省" in str(h))), None)
            continue
        province = ""
        if idx_prov is not None and idx_prov < len(cells) and cells[idx_prov] is not None:
            province = str(cells[idx_prov]).strip()
        rows.append({"row": n, "path": text, "province": province})
    if not rows:
        raise HTTPException(400, "Excel 中没有可导入的部门路径（请在第 1 列填部门级联路径）")
    return rows


# ---------- 接口 ----------

@router.get("")
def list_departments(db: Session = Depends(get_db)):
    """读取部门层级树（根 -> 子），供级联选择 / 管理展示。"""
    return dept.load_tree(db)


@router.get("/stats")
def department_stats(db: Session = Depends(get_db)):
    return {"total": dept.node_count(db)}


@router.get("/import/template")
def download_import_template():
    filename = quote("组织架构导入模板.xlsx")
    return Response(
        content=_build_template_bytes(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{filename}"},
    )


@router.post("/import")
def import_departments(file: UploadFile, db: Session = Depends(get_db)):
    """Excel 批量导入部门层级。逐行容错；同路径重复导入不重复建。"""
    suffix = Path(file.filename or "").suffix.lower()
    if suffix not in XLSX_SUFFIX:
        raise HTTPException(400, "仅支持 .xlsx / .xlsm 格式（可先下载模板）")
    rows = _read_import_rows(file.file)
    return dept.import_paths(db, rows)


@router.post("/reset")
def reset_default(db: Session = Depends(get_db)):
    """清空全部部门并恢复为内置默认示例树。"""
    dept.reset_to_default(db)
    return {"ok": True, "total": dept.node_count(db)}


@router.delete("/{department_id}")
def delete_department(department_id: int, db: Session = Depends(get_db)):
    """删除部门节点；存在下级部门时拒绝。"""
    try:
        dept.delete_node(db, department_id)
    except ValueError as e:
        raise HTTPException(400, str(e)) from e
    return {"ok": True}


@router.patch("/{department_id}")
def update_department(department_id: int, body: DepartmentProvinceUpdate,
                      db: Session = Depends(get_db)):
    """设置部门节点的所在省份（空串 = 恢复为继承上级）。岗位省份由此推导。"""
    try:
        node = dept.set_province(db, department_id, body.province)
    except ValueError as e:
        raise HTTPException(404, str(e)) from e
    return {"ok": True, "id": node.id, "province": node.province}
