"""导出服务：Excel 候选名单，以及「名单 + 简历原件」打包 ZIP（openpyxl + zipfile）。

数据源统一为 match_results 表；简历原件路径取记录里的 resume_path 快照
（旧记录快照为空时回退到 resumes.file_path），原件已丢失的记录在名单里标注"原件缺失"。
"""
from __future__ import annotations

import io
import logging
import re
import zipfile
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy.orm import Session

from ..models import MatchResult

logger = logging.getLogger(__name__)

_HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
_HEADER_FONT = Font(color="FFFFFF", bold=True)

# 简历原件在 ZIP 内的目录名
RESUME_DIR_IN_ZIP = "简历原件"
MISSING_RESUME = "（原件缺失）"

# Windows/zip 不允许的文件名字符
_ILLEGAL_CHARS = re.compile(r'[\\/:*?"<>|\r\n\t]+')


def _records(db: Session, job_id: int) -> list[MatchResult]:
    return (
        db.query(MatchResult)
        .filter(MatchResult.job_id == job_id)
        .order_by(MatchResult.score.desc())
        .all()
    )


def _display_name(rec: MatchResult) -> str:
    """与前端一致的展示姓名：优先解析/修正后的 structured.name，回退文件名。"""
    if rec.resume is None:
        return f"简历{rec.resume_id}"
    structured = rec.resume.structured or {}
    return str(structured.get("name") or rec.resume.name or f"简历{rec.resume_id}")


def _resume_path(rec: MatchResult) -> str:
    """简历原件完整路径：优先用匹配记录里的路径快照，其次回退到简历表当前路径。"""
    if rec.resume_path:
        return rec.resume_path
    return (rec.resume.file_path if rec.resume else "") or ""


def _pack_basename(rec: MatchResult) -> str:
    """ZIP 内展示用基名：姓名_专业，去掉路径非法字符。"""
    structured = (rec.resume.structured if rec.resume else None) or {}
    major = str(structured.get("major") or "")
    base = f"{_display_name(rec)}_{major}" if major else _display_name(rec)
    return _ILLEGAL_CHARS.sub("_", base).strip("_ .") or f"简历{rec.resume_id}"


def _plan_resume_files(records: list[MatchResult]) -> list[tuple[MatchResult, Path, str]]:
    """规划打包清单：[(匹配记录, 原件路径, ZIP 内相对路径)]。

    原件不存在（被手工删除 / 磁盘丢失）的记录直接跳过，由名单列标注缺失；
    同名简历自动加序号，避免 ZIP 内互相覆盖。
    """
    plan: list[tuple[MatchResult, Path, str]] = []
    used: set[str] = set()
    for rec in records:
        raw = _resume_path(rec)
        src = Path(raw) if raw else None
        if src is None or not src.is_file():
            logger.warning("匹配记录 %s 的简历原件不可用，跳过打包: %s", rec.id, raw or "(空路径)")
            continue
        suffix = src.suffix or ".pdf"
        base = _pack_basename(rec)
        arcname = f"{RESUME_DIR_IN_ZIP}/{base}{suffix}"
        n = 2
        while arcname.lower() in used:
            arcname = f"{RESUME_DIR_IN_ZIP}/{base}_{n}{suffix}"
            n += 1
        used.add(arcname.lower())
        plan.append((rec, src, arcname))
    return plan


def _build_workbook(records: list[MatchResult], file_names: dict[int, str]) -> Workbook:
    """构建候选名单工作簿；file_names 为 匹配记录id -> ZIP 内简历文件名（缺失则不填）。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "候选名单"
    headers = ["序号", "姓名", "学历", "专业", "年龄", "研究方向", "技能",
               "综合分", "向量相似度", "匹配理由", "推送状态", "简历文件"]
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT

    for i, rec in enumerate(records, 1):
        s = (rec.resume.structured if rec.resume else None) or {}
        ws.append([
            i, _display_name(rec), s.get("education"), s.get("major"), s.get("age"),
            s.get("research"), "、".join(s.get("skills") or []),
            rec.score, rec.vector_score, rec.reason, rec.push_status,
            file_names.get(rec.id) or MISSING_RESUME,
        ])

    widths = [6, 10, 8, 16, 6, 30, 24, 8, 10, 50, 10, 34]
    for idx, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + idx)].width = w
    return wb


def _to_bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_match_list(db: Session, job_id: int) -> bytes:
    """仅导出候选名单 xlsx（含"简历文件"列，标注原件是否齐备）。"""
    records = _records(db, job_id)
    names = {rec.id: arcname for rec, _, arcname in _plan_resume_files(records)}
    return _to_bytes(_build_workbook(records, names))


def export_match_bundle(db: Session, job_id: int, job_title: str) -> bytes:
    """导出「候选名单 Excel + 简历原件」ZIP 包。"""
    records = _records(db, job_id)
    plan = _plan_resume_files(records)
    names = {rec.id: arcname for rec, _, arcname in plan}

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr(export_filename(job_title), _to_bytes(_build_workbook(records, names)))
        for _, src, arcname in plan:
            zf.write(src, arcname)  # 中文文件名由 zipfile 自动打 UTF-8 标记
    logger.info("岗位 %s 导出打包：名单 %d 条，简历原件 %d 份", job_id, len(records), len(plan))
    return buf.getvalue()


def _timestamped(job_title: str, ext: str) -> str:
    # 岗位名可能含 / : * 等字符：直接进文件名会让 ZIP 内多出一层目录或浏览器保存失败
    safe_title = _ILLEGAL_CHARS.sub("_", job_title or "").strip("_ .") or "岗位"
    ts = datetime.now().strftime("%Y%m%d_%H%M")
    return f"候选名单_{safe_title}_{ts}{ext}"


def export_filename(job_title: str) -> str:
    return _timestamped(job_title, ".xlsx")


def bundle_filename(job_title: str) -> str:
    return _timestamped(job_title, ".zip")
